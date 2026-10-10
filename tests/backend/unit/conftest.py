import io
from pathlib import Path

import pytest
from docx import Document

from app.ai.domain.models import Utterance

FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "transcripts"


def fixture_bytes(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def utt(text: str, i: int = 0) -> Utterance:
    return Utterance(id=f"u{i:04d}", seq=i, speaker="Lan", start=float(i), text=text)


@pytest.fixture
def docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("[00:00:05] Lan: Xin chao")
    doc.add_paragraph("[00:00:20] Minh: Chao chi")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
